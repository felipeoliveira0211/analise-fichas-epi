"""Project only CA/equipment/approved protection/expiry from the provided workbook.
Never executes workbook formulas and never exports manufacturer/person identifiers.
"""
import argparse,collections,datetime,gzip,hashlib,json,pathlib,re
import openpyxl

def project(rows):
    registry={}; counts=collections.Counter()
    for row in rows:
        if not any(v is not None for v in row): continue
        counts['inputRows']+=1
        raw=row[0]
        if not isinstance(raw,(int,float)) or not float(raw).is_integer() or raw<=0:
            raise ValueError('Invalid CA on source row '+str(counts['inputRows']+1))
        ca=str(int(raw)); entry=registry.setdefault(ca,{'names':set(),'dates':set(),'protections':set(),'suspect':0})
        if not isinstance(row[1],datetime.datetime): raise ValueError('Invalid expiry for CA '+ca)
        entry['dates'].add(row[1].strftime('%d/%m/%Y'))
        nature=str(row[6] or '').strip().lower()
        front_aligned=nature in ('nacional','importado','importada')
        if front_aligned and isinstance(row[7],str) and row[7].strip(): entry['names'].add(row[7].strip())
        lab=row[15]
        lab_aligned=lab is None or isinstance(lab,(int,float)) or str(lab).strip().isdigit()
        approved=str(row[12] or '').strip()
        approval_aligned=bool(re.match(r'^PROTE[ÇC][ÃA]O\b',approved,re.I))
        if front_aligned and lab_aligned and approval_aligned:
            entry['protections'].add(approved)
        else:
            entry['suspect']+=1;counts['rowsRequiringAlignmentReview']+=1
    records=[]
    for ca,entry in sorted(registry.items(),key=lambda p:int(p[0])):
        name=' / '.join(sorted(entry['names'])) or 'Nome do EPI indisponível: revisar arquivo original'
        protection='\n'.join(sorted(entry['protections']))
        if not protection: counts['casWithoutReliableProtection']+=1
        if len(entry['dates'])!=1: counts['casWithConflictingExpiry']+=1
        if entry['suspect']: counts['casWithSourceAlignmentWarnings']+=1
        if len(entry['names'])>1: counts['casWithMultipleNames']+=1
        if len(entry['protections'])>1: counts['casWithMultipleApprovedTexts']+=1
        records.append([ca,name,protection,next(iter(entry['dates'])) if len(entry['dates'])==1 else ''])
    counts['uniqueCas']=len(records);counts['duplicateRowsConsolidated']=counts['inputRows']-len(records)
    return records,dict(counts)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('workbook');args=parser.parse_args()
    source=pathlib.Path(args.workbook);digest=hashlib.sha256(source.read_bytes()).hexdigest()
    workbook=openpyxl.load_workbook(source,read_only=True,data_only=True)
    sheet=workbook.active;header=next(sheet.iter_rows(min_row=1,max_row=1,max_col=19,values_only=True))
    if header[0]!='#NRRegistroCA' or header[7]!='NomeEquipamento' or header[12]!='AprovadoParaLaudo':raise ValueError('Unexpected workbook schema')
    records,stats=project(sheet.iter_rows(min_row=2,max_col=19,values_only=True))
    destination=pathlib.Path(__file__).resolve().parent.parent/'data/default-ca';destination.mkdir(parents=True,exist_ok=True)
    chunks=[]
    for index,start in enumerate(range(0,len(records),500)):
        raw=json.dumps(records[start:start+500],ensure_ascii=False,separators=(',',':')).encode()
        blob=gzip.compress(raw,mtime=0);name=f'chunk-{index:03d}.json.gz';(destination/name).write_bytes(blob)
        chunks.append({'file':name,'records':len(records[start:start+500]),'sha256':hashlib.sha256(blob).hexdigest()})
    modified=workbook.properties.modified
    manifest={'version':digest[:20],'fileName':source.name,'fileSha256':digest,'source':'Arquivo fornecido pelo usuário: '+source.name+' (origem online ainda não confirmada)','referenceDate':'','fileModifiedAt':modified.isoformat() if modified else None,'referenceNote':'Data de referência oficial não informada. Data de modificação interna do XLSX não comprova atualização da base.','statistics':stats,'chunks':chunks}
    (destination/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'statistics':stats,'compressedBytes':sum((destination/c['file']).stat().st_size for c in chunks),'chunks':len(chunks),'fileModifiedAt':manifest['fileModifiedAt']},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
